"""Build the user documentation as a native, Word-compatible DOCX file."""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Anwenderdokumentation.md"
TARGET = ROOT / "Anwenderdokumentation.docx"

INLINE_TOKEN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)"
)
IMAGE_LINE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
HEADING_LINE = re.compile(r"^(#{1,6})\s+(.+)$")
NUMBERED_LINE = re.compile(r"^\d+\.\s+(.+)$")
CHECKBOX_LINE = re.compile(r"^-\s+\[([ xX])\]\s+(.+)$")
BULLET_LINE = re.compile(r"^[-*]\s+(.+)$")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def clean_inline_text(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").replace("*", "")


def add_inline_runs(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(92, 56, 21)
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            paragraph.add_run(link_match.group(1) if link_match else token)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_text_paragraph(document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    return paragraph


def parse_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(document, lines: list[str]) -> None:
    parsed_rows = [parse_table_row(line) for line in lines]
    parsed_rows = [row for row in parsed_rows if not is_table_separator(row)]
    if not parsed_rows:
        return
    column_count = max(len(row) for row in parsed_rows)
    table = document.add_table(rows=len(parsed_rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(parsed_rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            add_inline_runs(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
        if row_index == 0:
            set_repeat_table_header(table.rows[row_index])
    document.add_paragraph()


def add_code_block(document, code_lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    set_paragraph_shading(paragraph, "F2F2F2")
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    heading_sizes = {1: 24, 2: 17, 3: 13, 4: 11}
    for level, size in heading_sizes.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(27, 78, 121)
        style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        styles[list_style_name].font.name = "Arial"
        styles[list_style_name].font.size = Pt(10.5)
        styles[list_style_name].paragraph_format.space_after = Pt(2)

    header = section.header.paragraphs[0]
    header.text = "IDX Viewer – Anwenderdokumentation"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)

    document.core_properties.title = "IDX Viewer – Anwenderdokumentation"
    document.core_properties.subject = "Praxisorientierte Bedienungsanleitung"
    document.core_properties.author = "Codex"
    document.core_properties.comments = (
        "Native DOCX-Erzeugung mit python-docx; enthält annotierte Klick-Screenshots."
    )


def build_document(markdown: str) -> Document:
    document = Document()
    configure_document(document)

    lines = markdown.splitlines()
    index = 0
    image_number = 0
    in_code = False
    code_lines: list[str] = []
    first_heading_seen = False

    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw_line)
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue

        image_match = IMAGE_LINE.match(line.strip())
        if image_match:
            alt_text, relative_path = image_match.groups()
            image_path = ROOT / relative_path
            if not image_path.exists():
                raise FileNotFoundError(f"Screenshot fehlt: {image_path}")
            image_number += 1
            picture_paragraph = document.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.keep_with_next = True
            picture_paragraph.add_run().add_picture(
                str(image_path), width=Inches(6.35)
            )
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(9)
            caption.paragraph_format.keep_together = True
            caption_run = caption.add_run(
                f"Abbildung {image_number}: {clean_inline_text(alt_text)}"
            )
            caption_run.italic = True
            caption_run.font.size = Pt(8.5)
            caption_run.font.color.rgb = RGBColor(89, 89, 89)
            index += 1
            continue

        heading_match = HEADING_LINE.match(line)
        if heading_match:
            markdown_level = len(heading_match.group(1))
            heading_text = clean_inline_text(heading_match.group(2))
            word_level = min(markdown_level, 4)
            if first_heading_seen and markdown_level == 2 and re.match(
                r"\d+\.", heading_text
            ):
                document.add_page_break()
            heading = document.add_heading(heading_text, level=word_level)
            if not first_heading_seen:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                heading.paragraph_format.space_after = Pt(8)
                first_heading_seen = True
            index += 1
            continue

        checkbox_match = CHECKBOX_LINE.match(line)
        if checkbox_match:
            checked = checkbox_match.group(1).lower() == "x"
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(
                paragraph, ("☒ " if checked else "☐ ") + checkbox_match.group(2)
            )
            index += 1
            continue

        bullet_match = BULLET_LINE.match(line)
        if bullet_match:
            add_text_paragraph(document, bullet_match.group(1), "List Bullet")
            index += 1
            continue

        numbered_match = NUMBERED_LINE.match(line)
        if numbered_match:
            add_text_paragraph(document, numbered_match.group(1), "List Number")
            index += 1
            continue

        if re.fullmatch(r"-{3,}", line.strip()):
            index += 1
            continue

        add_text_paragraph(document, line)
        index += 1

    if in_code:
        add_code_block(document, code_lines)

    return document


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    document = build_document(markdown)
    with tempfile.NamedTemporaryFile(
        prefix="Anwenderdokumentation_", suffix=".docx", dir=ROOT, delete=False
    ) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        document.save(temp_path)
        temp_path.replace(TARGET)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    print(
        f"Erzeugt: {TARGET} | "
        f"Absätze: {len(document.paragraphs)} | "
        f"Tabellen: {len(document.tables)} | "
        f"Abbildungen: {len(document.inline_shapes)}"
    )


if __name__ == "__main__":
    main()
